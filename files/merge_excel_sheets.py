#!/usr/bin/env python3
"""Merge the same worksheet from multiple .xlsx/.xlsm files into one .xlsx."""

import argparse
from pathlib import Path
from openpyxl import Workbook, load_workbook

EXCEL_MAX_ROWS = 1_048_576


def merge_excel_files(files: list[Path], sheet_name: str, output: Path) -> int:
    expected_header = None
    total_rows = 0
    output_wb = Workbook(write_only=True)
    output_sheet = output_wb.create_sheet("Merged")

    for file_path in files:
        if not file_path.is_absolute():
            raise ValueError(f"Input path must be absolute: {file_path}")
        if not file_path.is_file():
            raise FileNotFoundError(f"Input file not found: {file_path}")
        if file_path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValueError(f"Unsupported input type: {file_path} (use .xlsx or .xlsm)")
        if file_path.resolve() == output.resolve():
            raise ValueError(f"Output file cannot also be an input: {file_path}")

        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Worksheet {sheet_name!r} missing in: {file_path}")
            source_sheet = workbook[sheet_name]
            rows = source_sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None or all(value is None for value in header):
                raise ValueError(f"Missing header in worksheet {sheet_name!r}: {file_path}")

            if expected_header is None:
                if "Source_File" in header:
                    raise ValueError("Input already contains a 'Source_File' column")
                expected_header = header
                output_sheet.append((*header, "Source_File"))
            elif header != expected_header:
                raise ValueError(
                    f"Header mismatch in {file_path}\n"
                    f"Expected: {expected_header}\nFound:    {header}"
                )

            count = 0
            for row in rows:
                if all(value is None for value in row):
                    continue  # Ignore entirely empty rows.
                if total_rows + 2 > EXCEL_MAX_ROWS:
                    raise ValueError("Too many records for one Excel worksheet (max 1,048,575 data rows)")
                output_sheet.append((*row, file_path.name))
                count += 1
                total_rows += 1
            print(f"Merged {count:,} rows from {file_path}")
        finally:
            workbook.close()

    output.parent.mkdir(parents=True, exist_ok=True)
    output_wb.save(output)
    return total_rows


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Merge one named worksheet from multiple Excel files, row by row."
    )
    parser.add_argument("--sheet", required=True, help="Worksheet name shared by input files")
    parser.add_argument("--output", required=True, type=Path, help="Absolute output .xlsx path")
    parser.add_argument("--files", required=True, nargs="+", type=Path, help="Absolute input paths")
    args = parser.parse_args()

    if not args.output.is_absolute() or args.output.suffix.lower() != ".xlsx":
        parser.error("--output must be an absolute path ending in .xlsx")
    try:
        count = merge_excel_files(args.files, args.sheet, args.output)
    except (ValueError, FileNotFoundError, PermissionError, OSError) as exc:
        parser.exit(1, f"Error: {exc}\n")
    print(f"Done: {count:,} records saved to {args.output}")


if __name__ == "__main__":
    main()
